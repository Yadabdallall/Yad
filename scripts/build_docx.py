# -*- coding: utf-8 -*-
"""دروستکردنی خشتەی چالاکییە وەرزشییەکان بە فۆرماتی Word (.docx).

فۆنتی هەموو دەقەکان: Times New Roman، ئاراستە: ڕاست بۆ چەپ.
"""

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

FONT = "Times New Roman"
HEADER_BG = "1E3A8A"
STRIPE_BG = "F1F5F9"
BORDER = "BFC7D1"

TITLE = "خشتەی چالاکییە وەرزشییەکان"
HEADERS = ["ژ", "ناوی چالاکی", "بەروار", "ژمارەی بەشداربوان"]
# ژ ٥٪ | ناوی چالاکی ٥٠٪ | بەروار ١٥٪ | ژمارەی بەشداربوان ٣٠٪  (کۆی گشتی ٩٩٠٦ twip)
COL_WIDTHS = [495, 4953, 1486, 2972]
CENTERED = {0, 2}  # ستوونی ژمارە و بەروار لە ناوەڕاست

# ڕیزبەندیی پێویستی سکیمای OOXML بۆ منداڵەکانی هەر توخمێک
ORDER = {
    "w:rPr": [
        "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
        "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
        "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
        "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position",
        "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
        "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
        "w:eastAsianLayout", "w:specVanish", "w:oMath",
    ],
    "w:pPr": [
        "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore",
        "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
        "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
        "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
        "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange",
    ],
    "w:tcPr": [
        "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge",
        "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
        "w:tcFitText", "w:vAlign", "w:hideMark",
    ],
    "w:tblPr": [
        "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
        "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc",
        "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
        "w:tblLayout", "w:tblCellMar", "w:tblLook",
    ],
    "w:trPr": [
        "w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore",
        "w:wAfter", "w:cantSplit", "w:trHeight", "w:tblHeader",
        "w:tblCellSpacing", "w:jc", "w:hidden",
    ],
    "w:sectPr": [
        "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz", "w:pgMar",
        "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType", "w:cols",
        "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
        "w:printerSettings", "w:sectPrChange",
    ],
}


def _tag(el):
    """ناوی توخم بە شێوەی 'w:xxx'."""
    return "w:" + el.tag.split("}", 1)[1]


def put(parent, tag, **attrs):
    """توخمێک دروست بکە و لە شوێنی دروستی سکیمادا دایبنێ (یان ئەوەی هەیە بگەڕێنەوە)."""
    existing = parent.find(qn(tag))
    if existing is not None:
        el = existing
    else:
        el = OxmlElement(tag)
        order = ORDER[_tag(parent)]
        index = order.index(tag)
        anchor = None
        for child in parent:
            child_tag = _tag(child)
            if child_tag in order and order.index(child_tag) > index:
                anchor = child
                break
        if anchor is None:
            parent.append(el)
        else:
            anchor.addprevious(el)
    for key, value in attrs.items():
        el.set(qn("w:" + key), str(value))
    return el


def style_paragraph(paragraph, align):
    pf = paragraph.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.15
    put(paragraph._p.get_or_add_pPr(), "w:bidi")


def add_run(paragraph, text, size=11, bold=False, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = put(rPr, "w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn("w:" + attr), FONT)
    # قەبارە و ڕەقی بۆ پیتی کۆمپلێکس (کوردی/عەرەبی) + ئاراستەی ڕاست بۆ چەپ
    put(rPr, "w:szCs", val=str(int(size * 2)))
    if bold:
        put(rPr, "w:bCs", val="1")
    put(rPr, "w:rtl")
    return run


def shade(cell, fill):
    put(cell._tc.get_or_add_tcPr(), "w:shd", val="clear", color="auto", fill=fill)


def set_table_props(table, widths):
    tblPr = table._tbl.tblPr
    put(tblPr, "w:bidiVisual")  # خشتە بە ئاراستەی ڕاست بۆ چەپ

    borders = put(tblPr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)

    put(tblPr, "w:tblLayout", type="fixed")
    put(tblPr, "w:tblW", w=str(sum(widths)), type="dxa")

    grid = table._tbl.find(qn("w:tblGrid"))
    for col, width in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(width))


def fill_cell(cell, text, width, align, size=11, bold=False, color=None):
    tcPr = cell._tc.get_or_add_tcPr()
    put(tcPr, "w:tcW", w=str(width), type="dxa")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    margins = put(tcPr, "w:tcMar")
    for side, value in (("top", 80), ("left", 120), ("bottom", 80), ("right", 120)):
        m = OxmlElement("w:" + side)
        m.set(qn("w:w"), str(value))
        m.set(qn("w:type"), "dxa")
        margins.append(m)

    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, align)
    add_run(paragraph, text, size=size, bold=bold, color=color)


ROWS = [
    ("١", "یاریی کۆتایی پاڵەوانێتیی تۆپی بالەی تیپە میللییەکان لە سلێمانی",
     "28/4/2026", "١٨ تیپ (٢٥٢ یاریزانی کوڕ)"),
    ("٢", "خولی دادوەری و دابەشکردنی بڕوانامەی پشتێنەی ڕەش بۆ یاریزانانی پێشکەوتووی کیۆکۆشین لە پارکی هەواری شار",
     "1/5/2026", "٤٤ یاریزان"),
    ("٣", "مەراسیمی ڕێزلێنان لە تیپەکانی یانەی پێشمەرگە (تەمەن ژێر ١٧ و ٢١ ساڵ) بۆ سەرکەوتنیان بۆ خولی ئەستێرەکانی عێراق",
     "5/5/2026", "دیاری نەکراوە (تیپەکانی ژێر ١٧ و ٢١ ساڵ)"),
    ("٤", "پاڵەوانێتیی تێنسی سەرزەوی بۆ یانەکانی سنوری پارێزگای سلێمانی",
     "8/5/2026", "٧ یانە (٧٠ بۆ ٨٠ کوڕ، ١٥ بۆ ٢٠ کچ)"),
    ("٥", "پاڵەوانێتیی کراوەی سنووکەر لە سلێمانی",
     "10/5/2026", "٦٤ یاریزان"),
    ("٦", "چالاکیی وەرزشیی تایبەت بە کێشە ئۆڵۆمپییەکانی تایکواندۆ لە فامیلی مۆڵ",
     "14/5/2026", "٢٩ یاریزان (١٨ کوڕ، ١١ کچ)"),
    ("٧", "پاڵەوانێتیی بەرزکردنەوەی قورسایی بۆ خوێندکارانی کۆلێژی پەروەردەی وەرزشیی زانکۆی سلێمانی",
     "16/5/2026", "١١ خوێندکار (٤ کوڕ، ٧ کچ)"),
    ("٨", "سیمیناری زانستی و هونەری لە بوارەکانی (تاولۆ و ساندا) بۆ ڕاهێنەرانی ووشو کۆنگ فو",
     "20/5/2026", "ئامار بەردەست نییە"),
    ("٩", "پاڵەوانێتیی کراوەی بدمینتۆن (تۆپی پەڕ) بۆ کچان و کوڕان",
     "23/5/2026", "ئامار بەردەست نییە"),
    ("١٠", "خولی فێرگەکانی تۆپی پێ بۆ کچان و پێشوازیکردن لە شاندی هەڵبژاردەی کچانی عێراق",
     "15/6/2026", "٩ تیپ (١٤٠ یاریزانی کچ)"),
    ("١١", "پاڵەوانێتیی کەوان و تیر بۆ یانەکانی پارێزگای سلێمانی (کوڕان و کچان)",
     "19/6/2026", "٨ یانە (سلێمانی، سیروانی نوێ، ئاشتی، پێشمەرگە، هەڵەبجە، ئافرۆدێت، خورماڵ، نەورۆز)"),
    ("١٢", "فیستیڤاڵی وەرزشیی هەمەجۆر بەبۆنەی ڕۆژی جیهانیی ئۆڵۆمپی",
     "23/6/2026", "نزیکەی ٣٠٠ بۆ ٣٥٠ یاریزان"),
    ("١٣", "پاڵەوانێتیی «پادڵ بۆرد»ـی عێراق لە سلێمانی",
     "25/6/2026", "نزیکەی ٢٨ یاریزان (کوڕ)"),
    ("١٤", "پاڵەوانێتیی کراوەی تۆپی مێز لە سلێمانی بۆ سەرجەم تەمەنەکان",
     "28/6/2026", "نزیکەی ٨٠ یاریزانی کوڕ و کچ"),
    ("١٥", "یارییەکانی قۆناغی دووەمی خولی نایابی تۆپی دەستی عێراق لە هۆڵی زانکۆی سلێمانی",
     "5/7/2026", "٧ یانە (نزیکەی ١٠٠ یاریزان)"),
    ("١٦", "خولی پەرەپێدان لە هونەری «تاولۆ» و خولی دادوەریی «ساندا» لە ووشو کۆنگ فو",
     "5/7/2026", "٧٩ بەشداربوو (٦٧ تاولۆ + ١٢ دادوەریی ساندا)"),
    ("١٧", "پاڵەوانێتیی «سوپەری عێراق» بۆ کیک بۆکسینگ لە کۆلێژی پەروەردەی وەرزشی",
     "22/7/2026", "٩٦ یاریزان (٨١ کوڕ، ١٥ کچ)"),
    ("١٨", "پاڵەوانێتیی یانەکانی عێراق بۆ تایکواندۆ لە سلێمانی",
     "26/7/2026", "٨٣ یانە"),
    ("١٩", "فیستیڤاڵی «نوێکردنەوەی تۆپی سەبەتە» بۆ یانە و فێرگەکان (تەمەن ژێر ١٢ ساڵ)",
     "30/7/2026", "نزیکەی ٥٠ یاریزان"),
    ("٢٠", "خولی پەرەپێدان لە هەردوو وەرزشی «پیلاتس و ئایرۆبیک»",
     "31/7/2026", "نزیکەی ٣٥ ڕاهێنەری هۆڵەکان (کچان)"),
    ("٢١", "فیستیڤاڵی «مینی باسکێت» لە هۆڵی کۆلێژی پەروەردەی وەرزشی",
     "31/7/2026", "٢٠٠ یاریزان (کوڕ و کچ)"),
    ("٢٢", "یاری دۆستانەی تۆپی سەبەتە لەنێوان هەڵبژاردەی شاری سنە و شاری سلێمانی",
     "6/8/2026", "نزیکەی ٢٨ یاریزان"),
    ("٢٣", "پاڵەوانێتیی وەرزشیی جیۆ جیتسۆ لە زانکۆی سلێمانی",
     "13/8/2026", "٩٠ بەشداربوو"),
    ("٢٤", "پاڵەوانێتیی مەلەوانگەکانی سلێمانی بۆ کچان و کوڕان",
     "14/8/2026", "٥٠ یاریزانی کوڕ و کچ"),
    ("٢٥", "کۆبوونەوەی هونەری بۆ دەستپێکی پاڵەوانێتیی تۆپی بالەی فێرگە وەرزشییەکان",
     "18/8/2026", "١٤ تیپ لە ٧ فێرگە (١٩٦ یاریزانی کچ)"),
    ("٢٦", "خولی ڕاهێنەرایەتی و ناوبژیوانی لە وەرزشی ووشو کۆنگ فو",
     "27/8/2026", "٩٠ بەشداربوو (هەردوو ڕەگەز)"),
    ("٢٧", "پاڵەوانێتیی کراوەی شەتڕەنجی سلێمانی",
     "27/8/2026", "٤٩ یاریزانی کوڕ و کچ"),
    ("٢٨", "خولی ناوبژیوانیی تایکواندۆ",
     "28/8/2026", "٣٠ بەشداربوو (کوڕ و کچ)"),
    ("٢٩", "پاڵەوانێتیی ووشو کۆنگ فو (تاولۆ و ساندا)",
     "28/8/2026", "١١٠ یاریزان (٥٠ تاولۆ، ٦٠ ساندا)"),
    ("٣٠", "چوارەمین فیستیڤاڵی کاراتێی یانەی پێشمەرگە",
     "28/8/2026", "٣٥٠ یاریزان"),
    ("٣١", "دەستپێکی پاڵەوانێتیی تۆپی بالەی فێرگە وەرزشییەکان",
     "2/9/2026", "١٤ تیپ"),
    ("٣٢", "پاڵەوانێتیی گۆڕەپان و مەیدانی یانەکانی کوردستان و سلێمانی",
     "5/9/2026", "ئامار بەردەست نییە"),
]


def build(path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    rFonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn("w:" + attr), FONT)

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    section = doc.sections[0]
    section.top_margin = Twips(900)
    section.bottom_margin = Twips(900)
    section.left_margin = Twips(1000)
    section.right_margin = Twips(1000)
    put(section._sectPr, "w:bidi")

    title = doc.add_paragraph()
    style_paragraph(title, WD_ALIGN_PARAGRAPH.CENTER)
    title.paragraph_format.space_after = Pt(10)
    add_run(title, TITLE, size=16, bold=True, color=RGBColor(0x1E, 0x3A, 0x8A))

    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    set_table_props(table, COL_WIDTHS)

    head = table.rows[0]
    trPr = head._tr.get_or_add_trPr()
    put(trPr, "w:cantSplit")
    put(trPr, "w:tblHeader")  # دووبارەکردنەوەی سەرێک لە هەر لاپەڕەیەک
    for i, text in enumerate(HEADERS):
        align = WD_ALIGN_PARAGRAPH.CENTER if i in CENTERED else WD_ALIGN_PARAGRAPH.RIGHT
        fill_cell(head.cells[i], text, COL_WIDTHS[i], align, bold=True,
                  color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(head.cells[i], HEADER_BG)

    for n, values in enumerate(ROWS):
        row = table.add_row()
        put(row._tr.get_or_add_trPr(), "w:cantSplit")
        for i, text in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in CENTERED else WD_ALIGN_PARAGRAPH.RIGHT
            fill_cell(row.cells[i], text, COL_WIDTHS[i], align)
            if n % 2 == 1:
                shade(row.cells[i], STRIPE_BG)

    doc.save(path)
    print("نووسرا: %s" % path)


if __name__ == "__main__":
    import sys
    build(sys.argv[1] if len(sys.argv) > 1 else "chalakiyekan.docx")
