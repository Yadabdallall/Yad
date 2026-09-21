#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
دروستکەری کتێبی «سەرۆک کۆماری یاریگا».

سەرچاوە: book/src/*.md   →   دەرئەنجام: build/serok_komari_yariga.docx

نووسینەکە بە زمانێکی نیشانەکردنی سادە نووسراوە (@H1، @IMG، @QUOTE …) کە
لێرەدا دەگۆڕدرێت بۆ بەڵگەنامەیەکی Word ـی ڕاست-بۆ-چەپ.
"""
import os
import re
import sys
import glob
import json

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docxkit import (  # noqa: E402
    ordered, TBLPR_ORDER, TCPR_ORDER, BORDER_ORDER,
    GREEN_DARKEST, GREEN_DARK, GREEN_MID, GREEN_SOFT, GREEN_TINT,
    GREEN_TINT_2, GOLD, INK, INK_SOFT,
    FONT_BODY, FONT_DISPLAY,
    para, style_run, set_bidi, shade, borders, page_break, spacer,
    add_field, bookmark, section_rtl, set_page_number_format,
    keep_with_next, keep_lines, tabstop_right, ar, _el, page_break_before,
    add_hyperlink,
)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "book", "src")
IMG = os.path.join(ROOT, "book", "images")
OUT = os.path.join(ROOT, "build")

TITLE = "سەرۆک کۆماری یاریگا"
SUBTITLE = "مێژووی سیاسیی وەرزش لە عێراق — لە سێبەری ستەمکارییەوە بۆ چەترە گەورەکەی مام جەلال"
AUTHOR = "سۆلینی قەرەخەرمان"

TEXT_WIDTH_CM = 11.2

# پێوەری ڕێکخستنی درێژی کتێبەکە — بۆ گەیشتن بە ژمارەیەکی دیاریکراوی لاپەڕە
BODY_AFTER = float(os.environ.get("BOOK_BODY_AFTER", "5.0"))
BODY_LINE = float(os.environ.get("BOOK_BODY_LINE", "1.38"))
BODY_SIZE = float(os.environ.get("BOOK_BODY_SIZE", "11.0"))
HEAD_GAP = float(os.environ.get("BOOK_HEAD_GAP", "0.55"))  # ڕێژەی بۆشایی سەردێڕەکان


# ═══════════════════════════════════════════════════ ڕێکخستنی لاپەڕە
def setup_section(sec, *, header=True, footer=True, mirror=True):
    sec.page_width = Cm(14.8)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(2.0)
    section_rtl(sec)
    if mirror:
        sectPr = sec._sectPr
        pgMar = sectPr.find(qn("w:pgMar"))
        if pgMar is not None:
            pgMar.set(qn("w:gutter"), "0")
    sec.different_first_page_header_footer = False
    return sec


def build_footer(sec):
    """ژمارەی پەڕە لە ناوەڕاستی ژێرەوە، لەگەڵ دوو هێڵی سەوز."""
    f = sec.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
    p.text = ""
    set_bidi(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run("— "), size=9, color=GREEN_SOFT, font=FONT_DISPLAY)
    add_field(p, " PAGE ", size=10, color=GREEN_DARK, font=FONT_DISPLAY, bold=True)
    style_run(p.add_run(" —"), size=9, color=GREEN_SOFT, font=FONT_DISPLAY)
    return f


def build_header(sec, text):
    h = sec.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0] if h.paragraphs else h.add_paragraph()
    p.text = ""
    set_bidi(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(text), size=8.5, color=INK_SOFT, font=FONT_DISPLAY)
    borders(p, bottom=(4, GREEN_TINT_2), space=3)
    return h


def blank_header(sec):
    h = sec.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0] if h.paragraphs else h.add_paragraph()
    p.text = ""
    set_bidi(p)
    return h


def blank_footer(sec):
    f = sec.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
    p.text = ""
    set_bidi(p)
    return f


# ═══════════════════════════════════════════════════════ پێکهاتەکان
class Builder:
    def __init__(self, doc, toc_pages=None, toc_slots=4):
        self.doc = doc
        self.toc_pages = toc_pages or {}     # anchor -> page number
        self.toc_slots = toc_slots
        self.toc_entries = []                # (level, title, anchor)
        self.bid = 1000
        self.tok = 0
        self.pending_break = False

    # ------------------------------------------------- پەڕەبڕی دواخراو
    def _para(self, *a, **kw):
        """بڕگەیەک دروست دەکات و پەڕەبڕی چاوەڕوان بەسەریدا جێبەجێ دەکات."""
        p = para(self.doc, *a, **kw)
        if self.pending_break:
            page_break_before(p)
            self.pending_break = False
        return p

    def _spacer(self, pts):
        p = spacer(self.doc, pts)
        if self.pending_break:
            page_break_before(p)
            self.pending_break = False
        return p

    # ---------------------------------------------------- نیشانەی پێڕست
    def _token(self, paragraph, anchor):
        """نیشانەیەکی بچووکی سپی بۆ دۆزینەوەی ژمارەی پەڕە لە PDF ـەکەدا.

        لە دەرئەنجامی کۆتاییدا لادەبرێت (BOOK_TOKENS=0).
        """
        if os.environ.get("BOOK_TOKENS", "1") == "0":
            return
        r = paragraph.add_run(f" @@{anchor}@@ ")
        style_run(r, size=1, color="FFFFFF", font=FONT_BODY, rtl=False)

    def _register(self, level, title, paragraph):
        anchor = f"T{len(self.toc_entries):03d}"
        self.toc_entries.append((level, title, anchor))
        self.bid += 1
        bookmark(paragraph, anchor, self.bid)
        self._token(paragraph, anchor)
        return anchor

    # ------------------------------------------------------- سەردێڕەکان
    def part(self, number, title, subtitle=""):
        """پەڕەیەکی تەواو بۆ دەستپێکی بەشێک."""
        self.pending_break = True
        self._spacer(110)

        p = self._para(align="center", after=0, line=1)
        style_run(p.add_run("◆"), size=13, color=GOLD, font=FONT_DISPLAY)

        p = self._para(align="center", before=10, after=4, line=1.1)
        style_run(p.add_run(f"بەشی {number}"), size=13, color=GREEN_SOFT,
                  font=FONT_DISPLAY)
        self._register(0, f"بەشی {number}: {title}", p)

        p = self._para(align="center", before=6, after=8, line=1.25)
        style_run(p.add_run(title), size=21, bold=True, color=GREEN_DARKEST,
                  font=FONT_DISPLAY)

        p = self._para(align="center", after=0, line=1)
        style_run(p.add_run("▬▬▬▬▬"), size=9, color=GOLD, font=FONT_DISPLAY)

        if subtitle:
            self._para(subtitle, align="center", size=11.5, italic=True,
                 color=INK_SOFT, before=12, after=0, line=1.5,
                 indent_r=0.8, indent_l=0.8)
        self.pending_break = True
        return self

    def h1(self, text):
        p = self._para(align="right", before=22 * HEAD_GAP, after=2, line=1.2)
        style_run(p.add_run(text), size=16.5, bold=True, color=GREEN_DARK,
                  font=FONT_DISPLAY)
        keep_with_next(p)
        keep_lines(p)
        self._register(1, text, p)
        rule = self._para(align="right", before=0, after=12 * HEAD_GAP, line=1)
        style_run(rule.add_run("▬▬▬▬"), size=7, color=GOLD, font=FONT_DISPLAY)
        keep_with_next(rule)
        return self

    def h2(self, text):
        p = self._para(align="right", before=16 * HEAD_GAP, after=6 * HEAD_GAP, line=1.25)
        style_run(p.add_run(text), size=13.5, bold=True, color=GREEN_MID,
                  font=FONT_DISPLAY)
        keep_with_next(p)
        keep_lines(p)
        self._register(2, text, p)
        return self

    def h3(self, text):
        p = self._para(align="right", before=12 * HEAD_GAP, after=4 * HEAD_GAP, line=1.25)
        style_run(p.add_run(text), size=12, bold=True, color=GREEN_SOFT,
                  font=FONT_BODY)
        keep_with_next(p)
        keep_lines(p)
        return self

    # --------------------------------------------------------- دەقەکان
    def body(self, text):
        self._para(text, align="both", size=BODY_SIZE, color=INK,
             before=0, after=BODY_AFTER, line=BODY_LINE, indent_first=0.5)
        return self

    def lead(self, text):
        p = self._para(text, align="both", size=12.5, color=GREEN_DARKEST,
                 before=4, after=12, line=1.45, indent_r=0.2, indent_l=0.2)
        return self

    def note(self, text):
        p = self._para(align="both", size=10, before=8, after=8,
                 line=1.4, indent_r=0.5, indent_l=0.5)
        style_run(p.add_run(text), size=10, italic=True, color=INK_SOFT,
                  font=FONT_BODY)
        return self

    def quote(self, text, attrib=""):
        p = self._para(align="both", before=12, after=2, line=1.45,
                 indent_r=0.55, indent_l=0.55)
        style_run(p.add_run("«" + text + "»"), size=11.5, italic=True,
                  color=GREEN_DARKEST, font=FONT_BODY)
        shade(p, GREEN_TINT)
        borders(p, right=(18, GREEN_MID), space=8)
        keep_lines(p)
        if attrib:
            keep_with_next(p)      # ناوی خاوەنی وتە لە وتەکەی جیا نابێتەوە
            a = self._para(align="left", before=0, after=12, line=1.2,
                     indent_r=0.55, indent_l=0.55)
            style_run(a.add_run("— " + attrib), size=9.5, bold=True,
                      color=GREEN_SOFT, font=FONT_DISPLAY)
            shade(a, GREEN_TINT)
            borders(a, right=(18, GREEN_MID), space=8)
            keep_lines(a)
        return self

    def box(self, title, lines):
        p = self._para(align="right", before=12, after=3, line=1.2,
                 indent_r=0.4, indent_l=0.4)
        style_run(p.add_run("▪ " + title), size=11.5, bold=True,
                  color=GREEN_DARKEST, font=FONT_DISPLAY)
        shade(p, GREEN_TINT)
        borders(p, top=(8, GREEN_MID), right=(8, GREEN_MID),
                left=(8, GREEN_MID), space=7)
        keep_with_next(p)
        for i, ln in enumerate(lines):
            last = (i == len(lines) - 1)
            q = self._para(align="both", before=0,
                     after=8 if last else 4, line=1.4,
                     indent_r=0.4, indent_l=0.4)
            style_run(q.add_run(ln), size=10.5, color=INK, font=FONT_BODY)
            shade(q, GREEN_TINT)
            borders(q, right=(8, GREEN_MID), left=(8, GREEN_MID),
                    bottom=(8, GREEN_MID) if last else None, space=7)
            if not last:
                keep_with_next(q)
        return self

    def sources(self, items, title="سەرچاوەکانی ئەم بەشە"):
        """بڵۆکی سەرچاوەکان لە کۆتایی هەر بەشێکدا."""
        h = self._para(align="right", before=14, after=3, line=1.2,
                       indent_r=0.3, indent_l=0.3)
        style_run(h.add_run("▪ " + title), size=9.5, bold=True,
                  color=GREEN_DARKEST, font=FONT_DISPLAY)
        shade(h, "F4F8F5")
        borders(h, top=(6, GREEN_SOFT), right=(6, GREEN_SOFT),
                left=(6, GREEN_SOFT), space=7)
        keep_with_next(h)
        keep_lines(h)

        for i, it in enumerate(items):
            last_item = (i == len(items) - 1)

            # بەستەر لە دەقەکە جیا دەکرێتەوە و لە دێڕێکی تایبەتدا دادەنرێت
            url = None
            m = re.search(r"(https?://\S+)", it)
            if m:
                url = m.group(1).rstrip(".،")
                it = it[:m.start()].rstrip(" —«»,،")

            last = last_item and not url
            q = self._para(align="both", before=0, after=6 if last else 2,
                           line=1.3, indent_r=0.75, indent_l=0.3)
            q.paragraph_format.first_line_indent = Cm(-0.42)
            style_run(q.add_run(f"{ar(i + 1)}. "), size=8.5, bold=True,
                      color=GREEN_SOFT, font=FONT_DISPLAY)
            style_run(q.add_run(it), size=8.5, color=INK_SOFT, font=FONT_BODY)
            shade(q, "F4F8F5")
            borders(q, right=(6, GREEN_SOFT), left=(6, GREEN_SOFT),
                    bottom=(6, GREEN_SOFT) if last else None, space=7)
            keep_lines(q)
            if not last:
                keep_with_next(q)

            if url:
                u = self._para(align="left", before=0,
                               after=6 if last_item else 3,
                               line=1.15, indent_r=1.1, indent_l=0.3)
                style_run(u.add_run("↗ "), size=7.5, color=GREEN_SOFT,
                          font=FONT_DISPLAY)
                add_hyperlink(u, url, size=7.5, color="1F5C3A")
                shade(u, "F4F8F5")
                borders(u, right=(6, GREEN_SOFT), left=(6, GREEN_SOFT),
                        bottom=(6, GREEN_SOFT) if last_item else None,
                        space=7)
                keep_lines(u)
                if not last_item:
                    keep_with_next(u)
        return self

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            marker = f"{ar(i)}. " if numbered else "▪  "
            p = self._para(align="both", before=0, after=5, line=1.45,
                     indent_r=0.75)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            style_run(p.add_run(marker), size=11, bold=True, color=GREEN_SOFT,
                      font=FONT_DISPLAY)
            style_run(p.add_run(it), size=11.5, color=INK, font=FONT_BODY)
        return self

    def hr(self):
        p = self._para(align="center", before=10, after=12, line=1)
        style_run(p.add_run("❖  ❖  ❖"), size=10, color=GOLD, font=FONT_DISPLAY)
        return self

    def sig(self, text):
        # واژۆ نابێت بە تەنها لە سەری پەڕەیەکدا بمێنێتەوە
        prev = self.doc.paragraphs[-1] if self.doc.paragraphs else None
        if prev is not None:
            keep_with_next(prev)
        p = self._para(align="left", before=10, after=14, line=1.3)
        style_run(p.add_run(text), size=10.5, bold=True, color=GREEN_SOFT,
                  font=FONT_DISPLAY)
        return self

    # ---------------------------------------------------------- خشتەکان
    def table(self, caption, rows):
        """خشتەیەکی سادە و ڕوون: سەرێکی سەوزی تۆخ، هێڵی ئاسۆیی نەرم،
        بێ هێڵی ستوونی، و بۆشاییەکی ئاسوودە لە ناو خانەکاندا."""
        if caption:
            c = self._para(align="right", before=16, after=5, line=1.2)
            style_run(c.add_run("◆  "), size=8, color=GOLD, font=FONT_DISPLAY)
            style_run(c.add_run(caption), size=10.5, bold=True,
                      color=GREEN_DARKEST, font=FONT_DISPLAY)
            keep_with_next(c)

        if self.pending_break:      # خشتەیەک بێ ناونیشان لە سەری پەڕەدا
            anchor = self._para(before=0, after=0, line=1)
            style_run(anchor.add_run(""), size=1)

        ncols = len(rows[0])
        t = self.doc.add_table(rows=0, cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False

        tblPr = t._tbl.tblPr
        ordered(tblPr, _el("w:bidiVisual"), TBLPR_ORDER)
        ordered(tblPr, _el("w:tblW", w=str(int(Cm(TEXT_WIDTH_CM).twips)),
                           type="dxa"), TBLPR_ORDER)

        # تەنها هێڵی ئاسۆیی — هیچ هێڵێکی ستوونی نییە
        bd = OxmlElement("w:tblBorders")
        ordered(bd, _el("w:top", val="single", sz="12", space="0",
                        color=GREEN_DARK), BORDER_ORDER)
        ordered(bd, _el("w:bottom", val="single", sz="12", space="0",
                        color=GREEN_DARK), BORDER_ORDER)
        ordered(bd, _el("w:insideH", val="single", sz="2", space="0",
                        color="C8DACE"), BORDER_ORDER)
        ordered(bd, _el("w:left", val="nil"), BORDER_ORDER)
        ordered(bd, _el("w:right", val="nil"), BORDER_ORDER)
        ordered(bd, _el("w:insideV", val="nil"), BORDER_ORDER)
        ordered(tblPr, bd, TBLPR_ORDER)

        # بۆشایی لە ناو خانەکاندا
        mar = OxmlElement("w:tblCellMar")
        for side, w in (("top", 60), ("bottom", 60),
                        ("left", 110), ("right", 110)):
            mar.append(_el("w:" + side, w=str(w), type="dxa"))
        ordered(tblPr, mar, TBLPR_ORDER)

        widths = [TEXT_WIDTH_CM / ncols] * ncols
        if ncols == 2:
            widths = [TEXT_WIDTH_CM * 0.32, TEXT_WIDTH_CM * 0.68]
        elif ncols == 3:
            widths = [TEXT_WIDTH_CM * 0.26, TEXT_WIDTH_CM * 0.32,
                      TEXT_WIDTH_CM * 0.42]

        for ri, row in enumerate(rows):
            trow = t.add_row()
            trPr = trow._tr.get_or_add_trPr()
            trPr.append(_el("w:cantSplit"))
            if ri == 0:
                trPr.append(_el("w:tblHeader"))
            cells = trow.cells

            for ci, val in enumerate(row):
                cell = cells[ci]
                cell.width = Cm(widths[ci])
                tcPr = cell._tc.get_or_add_tcPr()
                ordered(tcPr, _el("w:vAlign", val="center"), TCPR_ORDER)

                p = cell.paragraphs[0]
                set_bidi(p)
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if ri == 0
                               else WD_ALIGN_PARAGRAPH.RIGHT)
                pf = p.paragraph_format
                pf.space_before = Pt(1)
                pf.space_after = Pt(1)
                pf.line_spacing = 1.3

                if ri == 0:
                    style_run(p.add_run(val), size=9.5, bold=True,
                              color="FFFFFF", font=FONT_DISPLAY)
                    ordered(tcPr, _el("w:shd", val="clear", color="auto",
                                      fill=GREEN_DARK), TCPR_ORDER)
                else:
                    # ستوونی یەکەم وەک کلیل — تۆختر و بە ڕەنگی سەوز
                    first = (ci == 0)
                    style_run(p.add_run(val), size=9.5, bold=first,
                              color=GREEN_DARKEST if first else INK,
                              font=FONT_BODY)
                    if ri % 2 == 0:
                        ordered(tcPr, _el("w:shd", val="clear", color="auto",
                                          fill="F5F9F6"), TCPR_ORDER)
        self._spacer(12)
        return self

    # ------------------------------------------------------------ وێنە
    def image(self, filename, caption, credit="", link=""):
        path = os.path.join(IMG, filename) if filename else None
        have = path and os.path.exists(path)

        if have:
            p = self.doc.add_paragraph()
            set_bidi(p)
            if self.pending_break:
                page_break_before(p)
                self.pending_break = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1
            keep_with_next(p)
            run = p.add_run()
            from PIL import Image as PILImage
            with PILImage.open(path) as im:
                w, h = im.size
            max_w, max_h = Cm(10.4), Cm(8.2)
            scale = min(max_w / w, max_h / h)
            run.add_picture(path, width=int(w * scale), height=int(h * scale))
        else:
            p = self._para(align="center", before=12, after=0, line=1.3,
                     indent_r=0.9, indent_l=0.9)
            style_run(p.add_run("⬚"), size=22, color="A9C3B3", font=FONT_DISPLAY)
            shade(p, "F5F9F6")
            borders(p, top=(6, "B9CFC0"), right=(6, "B9CFC0"),
                    left=(6, "B9CFC0"), space=10)
            keep_with_next(p)
            q = self._para(align="center", before=0, after=0, line=1.35,
                     indent_r=0.9, indent_l=0.9)
            style_run(q.add_run("شوێنی وێنە — ئەم فایلە دابنێ:"), size=8.5,
                      color=INK_SOFT, font=FONT_DISPLAY)
            shade(q, "F5F9F6")
            borders(q, right=(6, "B9CFC0"), left=(6, "B9CFC0"), space=10)
            keep_with_next(q)
            r = self._para(align="center", before=0, after=3, line=1.35,
                     indent_r=0.9, indent_l=0.9)
            style_run(r.add_run(f"book/images/{filename}"), size=8.5,
                      bold=True, color=GREEN_SOFT, font=FONT_DISPLAY, rtl=False)
            shade(r, "F5F9F6")
            borders(r, right=(6, "B9CFC0"), left=(6, "B9CFC0"),
                    bottom=(6, "B9CFC0"), space=10)
            keep_with_next(r)

        cap = self._para(align="center", before=2, after=14, line=1.35,
                   indent_r=0.5, indent_l=0.5)
        style_run(cap.add_run("وێنە: "), size=9, bold=True, color=GREEN_SOFT,
                  font=FONT_DISPLAY)
        style_run(cap.add_run(caption), size=9, italic=True, color=INK_SOFT,
                  font=FONT_BODY)
        if credit:
            style_run(cap.add_run("  \u2066(" + credit + ")\u2069"), size=8,
                      color="7C8C82", font=FONT_BODY, rtl=False)
        if link:
            cap.paragraph_format.space_after = Pt(2)
            ln = self._para(align="center", before=0, after=14, line=1.15,
                            indent_r=0.4, indent_l=0.4)
            style_run(ln.add_run("بۆ دۆزینەوەی وێنەکە: "), size=7.5,
                      color=GREEN_SOFT, font=FONT_DISPLAY)
            add_hyperlink(ln, link, size=7.5, color="1F5C3A")
            borders(ln, bottom=(4, GREEN_TINT_2), space=6)
        else:
            borders(cap, bottom=(4, GREEN_TINT_2), space=6)
        return self

    def pb(self):
        # بۆشاییەکی هەڵواسراو پێش پەڕەبڕ، پەڕەیەکی بەتاڵ دروست دەکات
        paras = self.doc.paragraphs
        if paras and not paras[-1].text.strip():
            el = paras[-1]._p
            el.getparent().remove(el)
        self.pending_break = True
        return self


# ═══════════════════════════════════════════════════ خوێندنەوەی سەرچاوە
def parse(builder, text):
    lines = text.split("\n")
    i = 0
    buf = []

    def flush():
        if buf:
            joined = " ".join(x.strip() for x in buf if x.strip())
            if joined:
                builder.body(joined)
            buf.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        if not s:
            flush()
            i += 1
            continue

        if s.startswith("#"):          # لێدوان — پشتگوێ دەخرێت
            i += 1
            continue

        m = re.match(r"^@(\w+)\s*(.*)$", s)
        if not m:
            buf.append(s)
            i += 1
            continue

        tag, rest = m.group(1).upper(), m.group(2).strip()
        flush()

        if tag == "PART":
            bits = [x.strip() for x in rest.split("|")]
            builder.part(bits[0], bits[1] if len(bits) > 1 else "",
                         bits[2] if len(bits) > 2 else "")
        elif tag == "H1":
            builder.h1(rest)
        elif tag == "H2":
            builder.h2(rest)
        elif tag == "H3":
            builder.h3(rest)
        elif tag == "LEAD":
            builder.lead(rest)
        elif tag == "NOTE":
            builder.note(rest)
        elif tag == "SIG":
            builder.sig(rest)
        elif tag == "HR":
            builder.hr()
        elif tag == "PB":
            builder.pb()
        elif tag == "QUOTE":
            bits = [x.strip() for x in rest.split("|")]
            builder.quote(bits[0], bits[1] if len(bits) > 1 else "")
        elif tag == "IMG":
            bits = [x.strip() for x in rest.split("|")]
            builder.image(bits[0], bits[1] if len(bits) > 1 else "",
                          bits[2] if len(bits) > 2 else "",
                          bits[3] if len(bits) > 3 else "")
        elif tag == "BOX":
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("@ENDBOX"):
                if lines[i].strip():
                    body.append(lines[i].strip())
                i += 1
            builder.box(rest, body)
        elif tag == "SRC":
            items = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("@ENDSRC"):
                if lines[i].strip():
                    items.append(re.sub(r"^[-•\d.\s]*", "", lines[i].strip()))
                i += 1
            if items:
                builder.sources(items, rest or "سەرچاوەکانی ئەم بەشە")
        elif tag in ("LIST", "NUM"):
            items = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("@END"):
                if lines[i].strip():
                    items.append(re.sub(r"^[-•]\s*", "", lines[i].strip()))
                i += 1
            builder.bullets(items, numbered=(tag == "NUM"))
        elif tag == "TABLE":
            rows = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("@ENDTABLE"):
                t = lines[i].strip()
                if t.startswith("|"):
                    rows.append([c.strip() for c in t.strip("|").split("|")])
                i += 1
            if rows:
                builder.table(rest, rows)
        i += 1

    flush()
    return builder


# ═══════════════════════════════════════════════════════ بەرگ و پێڕست
def cover(doc):
    sec = doc.sections[0]
    setup_section(sec)
    blank_footer(sec)
    blank_header(sec)

    spacer(doc, 46)
    p = para(doc, align="center", after=0, line=1)
    style_run(p.add_run("◆ ◆ ◆"), size=12, color=GOLD, font=FONT_DISPLAY)

    p = para(doc, align="center", before=22, after=0, line=1)
    style_run(p.add_run("▬▬▬▬▬▬▬▬▬▬▬▬"), size=10, color=GREEN_MID,
              font=FONT_DISPLAY)

    p = para(doc, align="center", before=18, after=6, line=1.15)
    style_run(p.add_run(TITLE), size=29, bold=True, color=GREEN_DARKEST,
              font=FONT_DISPLAY)

    p = para(doc, align="center", before=2, after=0, line=1)
    style_run(p.add_run("▬▬▬▬▬▬▬▬▬▬▬▬"), size=10, color=GREEN_MID,
              font=FONT_DISPLAY)

    spacer(doc, 22)
    p = para(doc, align="center", before=0, after=4, line=1.4,
             indent_r=0.6, indent_l=0.6)
    style_run(p.add_run("مێژووی سیاسیی وەرزش لە عێراق"), size=13.5,
              color=GREEN_DARK, font=FONT_DISPLAY)

    p = para(doc, align="center", before=0, after=0, line=1.45,
             indent_r=0.5, indent_l=0.5)
    style_run(p.add_run("لە سێبەری ستەمکارییەوە بۆ چەترە گەورەکەی مام جەلال"),
              size=12, italic=True, color=INK_SOFT, font=FONT_BODY)

    spacer(doc, 118)
    p = para(doc, align="center", before=0, after=3, line=1.2)
    style_run(p.add_run("نووسینی"), size=10, color=INK_SOFT, font=FONT_DISPLAY)
    p = para(doc, align="center", before=0, after=0, line=1.2)
    style_run(p.add_run(AUTHOR), size=16, bold=True, color=GREEN_DARKEST,
              font=FONT_DISPLAY)

    spacer(doc, 30)
    p = para(doc, align="center", before=0, after=0, line=1)
    style_run(p.add_run("❖"), size=11, color=GOLD, font=FONT_DISPLAY)
    page_break(doc)


def colophon(doc):
    spacer(doc, 150)
    for txt, sz, bold, col in [
        (TITLE, 13, True, GREEN_DARKEST),
        (SUBTITLE, 9.5, False, INK_SOFT),
        ("", 9, False, INK),
        ("نووسینی: " + AUTHOR, 10, True, INK),
        ("چاپی یەکەم", 9.5, False, INK_SOFT),
        ("", 9, False, INK),
        ("هەموو مافەکانی پارێزراون بۆ نووسەر.", 9, False, INK_SOFT),
        ("ئەم بەرهەمە پێشکەشە بە بۆردی وەرزشی یەکێتیی نیشتمانیی کوردستان.",
         9, False, INK_SOFT),
    ]:
        if not txt:
            spacer(doc, 8)
            continue
        para(doc, txt, align="center", size=sz, bold=bold, color=col,
             before=0, after=5, line=1.45, indent_r=0.4, indent_l=0.4,
             font=FONT_DISPLAY if bold else FONT_BODY)
    page_break(doc)


def dedication(doc):
    spacer(doc, 150)
    p = para(doc, align="center", before=0, after=0, line=1)
    style_run(p.add_run("❖"), size=13, color=GOLD, font=FONT_DISPLAY)
    spacer(doc, 16)
    para(doc, "پێشکەش", align="center", size=13, bold=True,
         color=GREEN_DARKEST, font=FONT_DISPLAY, before=0, after=14, line=1.2)
    for t in [
        "بە ڕۆحی هەموو ئەو وەرزشوانانەی عێراق کە لەژێر قامچی و ترسدا "
        "مەشقیان کرد و ناویان لە هیچ تابلۆیەکدا نەنووسرایەوە؛",
        "بە هەموو ئەو دایک و باوکانەی چاوەڕوانی کوڕەکانیان بوون لە "
        "دەرگای زەیوونەوە؛",
        "و بە یادی جەنابی مام جەلال، ئەو پیاوەی فێری کردین کە "
        "«ڕێزگرتن» خۆی بە تەنها خەڵاتێکی گەورەیە.",
    ]:
        para(doc, t, align="center", size=11, italic=True, color=INK_SOFT,
             before=0, after=11, line=1.6, indent_r=0.7, indent_l=0.7)
    page_break(doc)


def toc(doc, builder, entries, pages, slots):
    """پێڕستی ناوەڕۆک — بە ژمارەی پەڕەی ڕاستەقینە."""
    p = para(doc, align="center", before=30, after=2, line=1.2)
    style_run(p.add_run("پێڕستی ناوەڕۆک"), size=20, bold=True,
              color=GREEN_DARKEST, font=FONT_DISPLAY)
    p = para(doc, align="center", before=0, after=16, line=1)
    style_run(p.add_run("▬▬▬▬▬▬"), size=8, color=GOLD, font=FONT_DISPLAY)

    for level, title, anchor in entries:
        if level >= 2:          # پێڕستێکی کورت — تەنها بەش و بابەتەکان
            continue
        pg = pages.get(anchor)
        label = ar(pg) if pg else "…"

        if level == 0:
            q = para(doc, align="right", before=14, after=5, line=1.3,
                     indent_r=0.0)
            shade(q, GREEN_TINT)
            borders(q, right=(14, GREEN_MID), space=6)
            tabstop_right(q, TEXT_WIDTH_CM - 1.0)
            style_run(q.add_run(title), size=11.5, bold=True,
                      color=GREEN_DARKEST, font=FONT_DISPLAY)
            style_run(q.add_run("\t"), size=11.5, color=GREEN_SOFT)
            style_run(q.add_run(label), size=11, bold=True, color=GREEN_DARK,
                      font=FONT_DISPLAY)
        elif level == 1:
            q = para(doc, align="right", before=4, after=2, line=1.3,
                     indent_r=0.35)
            tabstop_right(q, TEXT_WIDTH_CM - 1.35)
            style_run(q.add_run(title), size=10.5, bold=True, color=GREEN_DARK,
                      font=FONT_BODY)
            style_run(q.add_run("\t"), size=10.5, color="9BB2A3")
            style_run(q.add_run(label), size=10, color=GREEN_SOFT,
                      font=FONT_DISPLAY)
        else:
            q = para(doc, align="right", before=1, after=1, line=1.25,
                     indent_r=0.85)
            tabstop_right(q, TEXT_WIDTH_CM - 1.85)
            style_run(q.add_run(title), size=9.5, color=INK, font=FONT_BODY)
            style_run(q.add_run("\t"), size=9.5, color="BFD0C6")
            style_run(q.add_run(label), size=9, color=INK_SOFT,
                      font=FONT_DISPLAY)


# ═══════════════════════════════════════════════════════════ سەرەکی
def source_files():
    return sorted(glob.glob(os.path.join(SRC, "*.md")))


def build(toc_pages=None, toc_slots=4, out_name="serok_komari_yariga.docx"):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT_BODY
    st.font.size = Pt(11.5)
    st.paragraph_format.space_after = Pt(6)

    # ---- بەرگ و ناونیشان: بێ سەرەوە و ژێرەوە
    cover(doc)
    colophon(doc)

    # ---- بەشێکی نوێ بۆ ئەوەی ژمارەی پەڕە لە بەرگدا دەرنەکەوێت
    body_sec = doc.add_section(WD_SECTION.ODD_PAGE
                               if False else WD_SECTION.NEW_PAGE)
    setup_section(body_sec)
    build_footer(body_sec)
    build_header(body_sec, TITLE)
    set_page_number_format(body_sec, fmt="hindiNumbers", start=3)

    dedication(doc)

    # ---- بەشی پێڕست (لە پاشدا پڕ دەکرێتەوە)
    toc_anchor_para = doc.add_paragraph()
    set_bidi(toc_anchor_para)

    b = Builder(doc, toc_pages=toc_pages, toc_slots=toc_slots)
    b.pending_break = True          # ناوەڕۆک لە پەڕەیەکی نوێوە دەست پێدەکات
    for f in source_files():
        with open(f, encoding="utf-8") as fh:
            parse(b, fh.read())

    # ---- پێڕست دەخرێتە شوێنی خۆی
    tmp = Document()
    tmp_sec = tmp.sections[0]
    setup_section(tmp_sec)
    toc(tmp, b, b.toc_entries, toc_pages or {}, toc_slots)
    toc_body = list(tmp.element.body)

    anchor_el = toc_anchor_para._p
    parent = anchor_el.getparent()
    idx = list(parent).index(anchor_el)
    for off, el in enumerate(x for x in toc_body if not x.tag.endswith("sectPr")):
        parent.insert(idx + off, el)

    parent.remove(anchor_el)        # بڕگەی بەتاڵی جێگرەوە لادەبرێت

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, out_name)
    doc.save(path)

    with open(os.path.join(OUT, "toc_entries.json"), "w", encoding="utf-8") as fh:
        json.dump(b.toc_entries, fh, ensure_ascii=False, indent=1)
    return path, b.toc_entries


if __name__ == "__main__":
    tp = {}
    if os.path.exists(os.path.join(OUT, "toc_pages.json")):
        with open(os.path.join(OUT, "toc_pages.json"), encoding="utf-8") as fh:
            tp = json.load(fh)
    path, entries = build(toc_pages=tp)
    print(f"دروستکرا: {path}  |  سەردێڕ: {len(entries)}")
