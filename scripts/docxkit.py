# -*- coding: utf-8 -*-
"""
docxkit — کیتی دروستکردنی بەڵگەنامەی ڕاست-بۆ-چەپ بۆ کوردیی سۆرانی.

OOXML پێویستی بە ڕیزبەندییەکی توندوتۆڵی ناوەکییە: هەر توخمێک دەبێت لە
شوێنی خۆیدا دابنرێت، ئەگەرنا Word و LibreOffice فایلەکە ڕەت دەکەنەوە.
`ordered()` ئەم ڕیزبەندییە بە خۆکارانە دەپارێزێت.
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

# ---------------------------------------------------------------- پاڵێت
GREEN_DARKEST = "0F3D26"
GREEN_DARK    = "14532D"
GREEN_MID     = "1E6B3C"
GREEN_SOFT    = "2D6A4F"
GREEN_TINT    = "E9F1EB"
GREEN_TINT_2  = "D6E6DB"
GOLD          = "9A7B2F"
INK           = "1A1A1A"
INK_SOFT      = "4A5A50"

# فۆنت — بە BOOK_FONT دەگۆڕدرێت، بۆ نموونە:
#   BOOK_FONT="Noto Naskh Arabic" python3 scripts/build.py
import os as _os
_F = _os.environ.get("BOOK_FONT", "Times New Roman")
FONT_BODY    = _F
FONT_DISPLAY = _F
FONT_LATIN   = _F

# ------------------------------------------- ڕیزبەندیی فەرمیی ECMA-376
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
    "pPrChange",
]

RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
    "dstrike", "outline", "shadow", "emboss", "imprint", "noProof",
    "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern",
    "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd",
    "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
    "specVanish", "oMath",
]

TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
    "tblDescription",
]

TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]

SECTPR_ORDER = [
    "footnotePr", "endnotePr", "type", "pgSz", "pgMar", "paperSrc",
    "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign",
    "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
    "printerSettings",
]

BORDER_ORDER = ["top", "left", "bottom", "right", "insideH", "insideV",
                "tl2br", "tr2bl"]


# ------------------------------------------------------- یاریدەدەری XML
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def ordered(parent, child, order, replace=True):
    """توخمێک لە شوێنی دروستی خۆیدا دادەنێت بەپێی ڕیزبەندیی OOXML."""
    name = child.tag.split("}")[-1]
    if name not in order:
        parent.append(child)
        return child
    rank = order.index(name)

    if replace:
        existing = parent.find(qn("w:" + name))
        if existing is not None:
            parent.remove(existing)

    for node in parent:
        nm = node.tag.split("}")[-1]
        if nm not in order or order.index(nm) > rank:
            node.addprevious(child)
            return child
    parent.append(child)
    return child


def _pPr(p):
    return p._p.get_or_add_pPr()


def _rPr(r):
    return r._r.get_or_add_rPr()


def add_pPr(p, tag, **attrs):
    return ordered(_pPr(p), _el("w:" + tag, **attrs), PPR_ORDER)


def add_rPr(r, tag, **attrs):
    return ordered(_rPr(r), _el("w:" + tag, **attrs), RPR_ORDER)


# --------------------------------------------------------- ڕازاندنەوە
def set_bidi(p):
    """ئاراستەی ڕاست-بۆ-چەپ بۆ بڕگەیەک."""
    add_pPr(p, "bidi")
    return p


def shade(p, fill):
    add_pPr(p, "shd", val="clear", color="auto", fill=fill)
    return p


def borders(p, right=None, left=None, top=None, bottom=None, space=6):
    """چوارچێوە. لە ڕاست-بۆ-چەپدا 'right' لای دەستپێکی دێڕە."""
    pBdr = OxmlElement("w:pBdr")
    for side, spec in (("top", top), ("left", left),
                       ("bottom", bottom), ("right", right)):
        if spec:
            sz, color = spec
            ordered(pBdr, _el("w:" + side, val="single", sz=sz,
                              space=space, color=color), BORDER_ORDER)
    ordered(_pPr(p), pBdr, PPR_ORDER)
    return p


def page_break_before(p):
    """پەڕەیەکی نوێ پێش ئەم بڕگەیە — بێ ئەوەی پەڕەی بەتاڵ دروست بکات."""
    add_pPr(p, "pageBreakBefore", val="1")
    return p


def keep_with_next(p):
    add_pPr(p, "keepNext", val="1")
    return p


def keep_lines(p):
    add_pPr(p, "keepLines", val="1")
    return p


def tabstop_right(p, pos_cm):
    """وێستانێک بە خاڵی پڕکەرەوە — بۆ پێڕستی ناوەڕۆک."""
    tabs = OxmlElement("w:tabs")
    tabs.append(_el("w:tab", val="left", leader="dot",
                    pos=int(Cm(pos_cm).twips)))
    ordered(_pPr(p), tabs, PPR_ORDER)
    return p


# ------------------------------------------------------------- ڕەوانەکان
def style_run(run, *, size=12, bold=False, italic=False, color=INK,
              font=FONT_BODY, spacing=None, rtl=True):
    """ڕەوانێک بە فۆنتی سکریپتی ئاڵۆز (cs) — بۆ کوردی پێویستە."""
    rPr = _rPr(run)

    rFonts = _el("w:rFonts", cs=font, ascii=FONT_LATIN, hAnsi=FONT_LATIN)
    ordered(rPr, rFonts, RPR_ORDER)

    run.font.size = Pt(size)
    ordered(rPr, _el("w:szCs", val=int(size * 2)), RPR_ORDER)

    run.font.bold = bold
    if bold:
        ordered(rPr, _el("w:bCs", val="1"), RPR_ORDER)
    run.font.italic = italic
    if italic:
        ordered(rPr, _el("w:iCs", val="1"), RPR_ORDER)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if spacing:
        ordered(rPr, _el("w:spacing", val=int(spacing)), RPR_ORDER)
    if rtl:
        ordered(rPr, _el("w:rtl"), RPR_ORDER)
    return run


def para(container, text="", *, align="both", size=12, bold=False,
         italic=False, color=INK, font=FONT_BODY, before=0, after=6,
         line=1.5, indent_first=0, indent_r=0, indent_l=0, style=None):
    """بڕگەیەکی ڕاست-بۆ-چەپ دروست دەکات و دەیڕازێنێتەوە."""
    p = container.add_paragraph(style=style)
    set_bidi(p)
    p.alignment = {
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
    if indent_first:
        pf.first_line_indent = Cm(indent_first)
    if indent_r:
        pf.right_indent = Cm(indent_r)
    if indent_l:
        pf.left_indent = Cm(indent_l)
    if text:
        style_run(p.add_run(text), size=size, bold=bold, italic=italic,
                  color=color, font=font)
    return p


# ------------------------------------------------------ پەڕە و خانەکان
def page_break(container):
    p = container.add_paragraph()
    set_bidi(p)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run()._r.append(_el("w:br", type="page"))
    return p


def spacer(container, pts):
    """بۆشاییەکی ستوونی بە خاڵ."""
    p = container.add_paragraph()
    set_bidi(p)
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(pts)
    style_run(p.add_run(" "), size=1)
    return p


def add_field(paragraph, instr, *, size=10, color=GREEN_DARK,
              font=FONT_DISPLAY, bold=False):
    """خانەیەکی خۆکار — وەک ژمارەی پەڕە."""
    def mk(child=None, text=None):
        r = paragraph.add_run(text or "")
        style_run(r, size=size, color=color, font=font, bold=bold)
        if child is not None:
            r._r.append(child)
        return r

    mk(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    mk(it)
    mk(_el("w:fldChar", fldCharType="separate"))
    mk(text="١")
    mk(_el("w:fldChar", fldCharType="end"))
    return paragraph


def bookmark(paragraph, name, bid):
    paragraph._p.insert(0, _el("w:bookmarkStart", id=bid, name=name))
    paragraph._p.append(_el("w:bookmarkEnd", id=bid))
    return paragraph


def section_rtl(section):
    ordered(section._sectPr, _el("w:bidi"), SECTPR_ORDER)
    return section


def set_page_number_format(section, fmt="decimal", start=None):
    attrs = {"fmt": fmt}
    if start is not None:
        attrs["start"] = start
    ordered(section._sectPr, _el("w:pgNumType", **attrs), SECTPR_ORDER)
    return section


# ------------------------------------------------------------- ژمارەکان
_AR_DIGITS = "٠١٢٣٤٥٦٧٨٩"


def ar(n):
    """ژمارەی لاتینی دەگۆڕێت بۆ ژمارەی عەرەبی-هیندی (٠-٩)."""
    return "".join(_AR_DIGITS[int(c)] if c.isdigit() else c for c in str(n))
